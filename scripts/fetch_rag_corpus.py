"""Fetch the open corpus used by the RAG evaluation suite (tests/test_rag.py).

The technical query set is evaluated against "Mastering Bitcoin" 3rd edition
(Antonopoulos & Harding), licensed CC BY-SA 4.0 (see its LICENSE file). The
text is NOT stored in this repository: this script downloads it from a pinned
commit of github.com/bitcoinbook/bitcoinbook, verifies every file's SHA-256,
and converts the AsciiDoc chapters into Markdown and into a DOCX (the upload
endpoint accepts PDF, PPTX and DOCX; headings become Word heading styles so the
ingestion pipeline keeps the chapter/section structure).

Usage (from services/ai):
    uv run python ../../scripts/fetch_rag_corpus.py [--out tests/data/corpus]

Output (git-ignored):
    <out>/mastering-bitcoin-3e.md   attribution header + all chapters
    <out>/mastering-bitcoin-3e.docx  same content, for upload
    <out>/LICENSE-mastering-bitcoin  the upstream licence file
"""
import argparse
import hashlib
import re
import sys
import urllib.request
from pathlib import Path

REPO = "bitcoinbook/bitcoinbook"
# Tag "third_edition_print1", pinned by commit so the corpus cannot drift.
COMMIT = "6d1c26e1640ae32b28389d5ae4caf1214c2be7db"

FILES = {
    "LICENSE": "92cc259d2e3af15ae8bd21bac5fab2121fb739a04a552e4ec68f78499fcbcd5e",
    "ch01_intro.adoc": "abd3aacb45ce188fb0abfc421b65360e967a008cb0480247042bc62f250bd310",
    "ch02_overview.adoc": "2fc3748e32c76c9ba87884b5c4e7173c55c72d465bb95448924c26049b83a769",
    "ch03_bitcoin-core.adoc": "fad1390d95bb33cb6ea9c7e4fd472eed7a1232051b47b8528ae4975403812496",
    "ch04_keys.adoc": "63e824f8dd77b32801c19e225fdb578f19d5b1db28954241f658f94e81235c11",
    "ch05_wallets.adoc": "3b113fedba6a9aae33c0c33dab1609b4a507049198a6aaa899bfb9c97783238a",
    "ch06_transactions.adoc": "74001f2cf450745c606f4aa06ef9e947818037cc03ae2b19d7c8f073cebdb3ed",
    "ch07_authorization-authentication.adoc": "ffe77488544b89b140174ef6911a93c7f3931b77673a83cc142fd5dd88850edb",
    "ch08_signatures.adoc": "d5308c9282503283fea4fea5d5dff1e6c2e1d72036b7858bad78a603d1dd02a7",
    "ch09_fees.adoc": "b5579a26064285c7dd5b7f1ccf7b27f3699f2251c3aad64028ef91a947c2a953",
    "ch10_network.adoc": "f2c43a33de5e92f814748aecc130c06e472111839a97e155efcdb9e5a8ee6107",
    "ch11_blockchain.adoc": "cabf040ec7b23714998381d340037c25ba248265202b5a4344590c3ff4464bd5",
    "ch12_mining.adoc": "987dc51ddafea7c08e5615fd30fd919902a5669ac807a48fb2039f46bc7a46e3",
    "ch13_security.adoc": "e1e401dc871bcb8ef1060232c3fd09c1854484f8c45284414e60f5537b061b90",
    "ch14_applications.adoc": "4643b17d93ec96f5a541b07f0c1e99a324d5cf73d53cf226519b51422de0a7b8",
}

CORPUS_NAME = "mastering-bitcoin-3e.md"
DOCX_NAME = "mastering-bitcoin-3e.docx"

ATTRIBUTION = f"""# Mastering Bitcoin, 3rd edition

Andreas M. Antonopoulos and David Harding. Licensed under the Creative Commons
Attribution-ShareAlike 4.0 International License
(https://creativecommons.org/licenses/by-sa/4.0/). Source:
https://github.com/{REPO} at commit {COMMIT} (tag third_edition_print1).
Converted from AsciiDoc to Markdown for retrieval evaluation; index markers,
anchors, images and cross-reference targets were removed.

"""

_INDEX = re.compile(r"\({3}.*?\){3}")
_ANCHOR = re.compile(r"\[\[[^\]]*\]\]")
_XREF = re.compile(r"<<([^,>]+)(?:,([^>]+))?>>")
_ATTR_LINE = re.compile(r"^\[[^\]]*\]$")
_PASS_INLINE = re.compile(r"pass:\[(.*?)\]")
_HTML_TAG = re.compile(r"<[^>]+>")
_MONO = re.compile(r"\+([^+\n]+)\+")


def asciidoc_to_markdown(text: str) -> str:
    """Convert the subset of AsciiDoc used by the book into plain Markdown."""
    out: list[str] = []
    in_code = False
    for line in text.splitlines():
        if line.strip() == "++++":
            # Passthrough blocks hold raw HTML (tables, lists): keep the text only.
            continue
        if line.strip() == "----":
            out.append("```")
            in_code = not in_code
            continue
        if in_code:
            out.append(line)
            continue
        if line.strip() in ("====", "****", "____") or line.startswith("image::"):
            continue
        if _ATTR_LINE.match(line.strip()):
            continue
        line = _PASS_INLINE.sub(lambda m: m.group(1), line)
        line = _HTML_TAG.sub(" ", line)
        line = _MONO.sub(lambda m: m.group(1), line)
        line = _INDEX.sub("", line)
        line = _ANCHOR.sub("", line)
        line = _XREF.sub(lambda m: (m.group(2) or "").strip(), line)
        heading = re.match(r"^(=+) (.*)$", line)
        if heading:
            line = "#" * len(heading.group(1)) + " " + heading.group(2)
        elif line.startswith(".") and len(line) > 1 and line[1].isalpha():
            line = "**" + line[1:] + "**"
        out.append(line.rstrip())
    return re.sub(r"\n{3,}", "\n\n", "\n".join(out)).strip() + "\n"


def markdown_to_docx(markdown: str, path: Path) -> None:
    """Write Markdown headings as Word heading styles, everything else as text."""
    from docx import Document  # backend dependency; run inside services/ai's env

    doc = Document()
    paragraph: list[str] = []

    def flush() -> None:
        if paragraph:
            doc.add_paragraph(" ".join(paragraph))
            paragraph.clear()

    in_code = False
    for line in markdown.splitlines():
        if line.startswith("```"):
            flush()
            in_code = not in_code
            continue
        heading = None if in_code else re.match(r"^(#{1,6}) (.*)$", line)
        if heading:
            flush()
            doc.add_heading(heading.group(2).strip(), level=min(len(heading.group(1)), 5))
        elif in_code:
            doc.add_paragraph(line)
        elif not line.strip():
            flush()
        else:
            paragraph.append(line.strip().replace("**", ""))
    flush()
    doc.save(str(path))


def fetch(name: str) -> bytes:
    url = f"https://raw.githubusercontent.com/{REPO}/{COMMIT}/{name}"
    with urllib.request.urlopen(url, timeout=60) as resp:
        data: bytes = resp.read()
    digest = hashlib.sha256(data).hexdigest()
    if digest != FILES[name]:
        raise SystemExit(f"Checksum mismatch for {name}: {digest} != {FILES[name]}")
    return data


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--out", default="tests/data/corpus", help="Output directory")
    args = parser.parse_args()
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)

    chapters = []
    for name in FILES:
        data = fetch(name)
        if name == "LICENSE":
            (out / "LICENSE-mastering-bitcoin").write_bytes(data)
        else:
            chapters.append(asciidoc_to_markdown(data.decode("utf-8")))

    corpus = out / CORPUS_NAME
    corpus.write_text(ATTRIBUTION + "\n".join(chapters), encoding="utf-8")
    markdown_to_docx(corpus.read_text(encoding="utf-8"), out / DOCX_NAME)
    print(f"Wrote {corpus} ({corpus.stat().st_size // 1024} KB, {len(chapters)} chapters)")
    print(f"Wrote {out / DOCX_NAME} ({(out / DOCX_NAME).stat().st_size // 1024} KB)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
